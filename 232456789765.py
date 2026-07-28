import asyncio
import datetime
import hashlib
import json
import os
import re
import subprocess
import time
from concurrent.futures import ThreadPoolExecutor, TimeoutError as FuturesTimeout
from io import BytesIO

import psycopg2
import requests
from docx import Document
from loguru import logger
from maxapi.types import ButtonsPayload, LinkButton
from yookassa import Payment, Configuration

from base_requests import send_xml_to_ekinobilet
from config import (
    db_path,
    info_log_file,
    kinopoisk_token,
    notify_admin,
    pochta_bank_token,
    root_path,
    send_document_sync,
    send_message_sync,
    url,
    url_kino_baza,
    url_prokultura,
    youkassa_shop_id,
    youkassa_secret_key,
    validation_url, url_server, get_loop,
)
from pkg.log import CustomLogger

async def check_payment_status():
    try:
        with psycopg2.connect(db_path) as conn:
            with conn.cursor() as curs:
                curs.execute(
                    "SELECT payment_id, payment_link, user_id, price, row, place, performance_id "
                    "FROM orders WHERE order_id = %s", (order_id,)
                )
                order = curs.fetchone()

                curs.execute(
                    "SELECT hallname, date, time, show_id, building_id "
                    "FROM performance WHERE performance_id = %s", (order[6],)
                )
                performance = curs.fetchone()

                curs.execute("SELECT name FROM show WHERE show_id = %s", (performance[3],))
                show = curs.fetchone()

                curs.execute("SELECT city FROM cinemas WHERE building_id = %s", (performance[4],))
                cinema = curs.fetchone()

                user_id = order[2]
                price = order[3]
                row = order[4]
                place = order[5]
                hallname = performance[0]
                date = performance[1]
                time_ = performance[2]
                name = show[0]
                city = cinema[0]

                months = {
                    1: "января", 2: "февраля", 3: "марта", 4: "апреля",
                    5: "мая", 6: "июня", 7: "июля", 8: "августа",
                    9: "сентября", 10: "октября", 11: "ноября", 12: "декабря",
                }
                date_obj = datetime.datetime.strptime(date, "%Y-%m-%d")
                date_ru = f"{date_obj.day} {months[date_obj.month]} {date_obj.year}"
                date_time_ru = f"{date_ru} {time_}"

                # Генерация PDF-билета
                template_path = os.path.join(root_path, "Шаблон билета.docx")
                doc = Document(template_path)

                now_plus7 = datetime.datetime.utcnow() + datetime.timedelta(hours=7)
                purchase_date_ru = (
                    f"{now_plus7.day} {months[now_plus7.month]} {now_plus7.year} "
                    f"{now_plus7.hour:02}:{now_plus7.minute:02}"
                )

                replace_dict = {
                    "{city}": city, "{cinema}": hallname,
                    "{number}": str(order_id), "{date}": purchase_date_ru,
                    "{cinema_date}": date_time_ru, "{name}": name,
                    "{row}": str(row), "{place}": str(place),
                    "{price}": str(price), "{fio}": fio,
                }

                def replace_placeholder_safe(paragraph, rd):
                    full_text = ''.join(r.text for r in paragraph.runs if r.text)
                    replaced = False
                    for key, value in rd.items():
                        if key in full_text:
                            full_text = full_text.replace(key, value)
                            replaced = True
                    if not replaced:
                        return
                    idx = 0
                    for run in paragraph.runs:
                        if run.text:
                            length = len(run.text)
                            run.text = full_text[idx:idx + length]
                            idx += length
                    if idx < len(full_text):
                        for run in reversed(paragraph.runs):
                            if run.text:
                                run.text += full_text[idx:]
                                break

                for paragraph in doc.paragraphs:
                    replace_placeholder_safe(paragraph, replace_dict)
                for table in doc.tables:
                    for trow in table.rows:
                        for cell in trow.cells:
                            for paragraph in cell.paragraphs:
                                replace_placeholder_safe(paragraph, replace_dict)

                docx_path = f"/tmp/Билет_№{order_id}.docx"
                pdf_path = f"/tmp/Билет_№{order_id}.pdf"
                doc.save(docx_path)

                subprocess.run(
                    ["libreoffice", "--headless", "--convert-to", "pdf",
                     docx_path, "--outdir", "/tmp"],
                    check=True,
                )

                msg_text = (
                    "✅ Заказ оплачен\n\n"
                    f"Заказ №: {order_id}\n"
                    f"Сеанс: {name}\n"
                    f"Кинотеатр: {hallname} ({city})\n"
                    f"Дата и время сеанса: {date_time_ru}\n"
                    f"Ряд / Место: {row} / {place}\n"
                    f"Цена: {price} р.\n\n"
                    "👇 Ваш билет ниже 👇"
                )
                await send_message_sync(chat_id, user_id, msg_text)
                await send_document_sync(
                    chat_id, user_id, pdf_path,
                    caption=f"Ваш билет №{order_id}",
                    filename=f"Билет_№{order_id}.pdf",
                )

                try:
                    os.remove(docx_path)
                    os.remove(pdf_path)
                except Exception:
                    logger.exception("Ошибка удаления временных файлов")

                # Добавляем оплату в МирКино
                params = {
                    "sp": "WgA_AddPayment",
                    "IdOrder": order_id,
                    "Amount": price,
                    "IdPaymentMethod": 11,
                    "idUser": 1,
                    "df": "J",
                }
                for _ in range(5):
                    try:
                        resp = requests.get(url_kino_baza, params=params)
                        break
                    except Exception:
                        time.sleep(7)
                payment_kino = resp.json()
                kino_add_payment_id = payment_kino['IdPayment']

                with psycopg2.connect(db_path) as conn2:
                    with conn2.cursor() as curs2:
                        try:
                            curs2.execute(
                                "UPDATE orders SET status = 1, kino_add_payment_id = %s "
                                "WHERE payment_id = %s",
                                (int(kino_add_payment_id), payment_id),
                            )
                        except (ValueError, TypeError):
                            curs2.execute(
                                "UPDATE orders SET status = 1 WHERE payment_id = %s",
                                (payment_id,),
                            )

                        curs2.execute(
                            "SELECT * FROM performance WHERE performance_id = %s",
                            (performance_id,)
                        )
                        performance_data = curs2.fetchone()

                        curs2.execute(
                            "SELECT fond_kino_id FROM cinemas WHERE building_id = %s",
                            (performance_data[2],)
                        )
                        fond_kino_id = curs2.fetchone()[0]

                        curs2.execute(
                            "SELECT pu_number, name, id_procult FROM show WHERE show_id = %s",
                            (performance_data[1],)
                        )
                        show_data = curs2.fetchone()

                        curs2.execute(
                            "SELECT report_sented FROM orders WHERE payment_id = %s", (payment_id,)
                        )
                        is_fk_report_send = curs2.fetchone()[0]

    except TypeError:
        err_msg = (
            f'!!!!Ошибка оформления заказа\n'
            f'order_id {order_data[0]}\nuser_id {order_data[1]}\n'
            f'performance {order_data[3]}\nplace_id {order_data[4]}\n'
            f'ряд {order_data[11]}\nместо {order_data[12]}\n'
            f'payment_id {order_data[6]}'
        )
        await notify_admin(err_msg)
        if report:
            await notify_admin(err_msg, admin_2=True)
        with psycopg2.connect(db_path) as conn:
            with conn.cursor() as curs:
                curs.execute(
                    "UPDATE orders SET status = 4 WHERE payment_id = %s", (payment_id,)
                )
    except Exception as e:
        logger.exception("Ошибка при обработке succeeded")
        await notify_admin(f'!!!!Ошибка оформления succeeded\n{e}')

    if not is_fk_report_send:
        await send_xml_to_ekinobilet(
            performance_data, show_data, payment, fond_kino_id,
            place, row, price, payment_id, order_id,
        )

    # Валидация Почта Банк
    try:
        rrn = payment.authorization_details.rrn
        with psycopg2.connect(db_path) as conn:
            with conn.cursor() as curs:
                curs.execute(
                    "SELECT date, time, show_id FROM performance WHERE performance_id = %s",
                    (performance_id,)
                )
                perf = curs.fetchone()
                curs.execute("SELECT id_procult FROM show WHERE show_id = %s", (perf[2],))
                event_id = curs.fetchone()[0]

        date_time_obj = datetime.datetime.strptime(f'{perf[0]} {perf[1]}', '%Y-%m-%d %H:%M')
        event_session_timestamp = int(time.mktime(date_time_obj.timetuple()))

        ok, text = validate_pochta_bank(
            buyer_info=fio, rrn=rrn, event_id=event_id,
            place_id=place_id, event_session_timestamp=event_session_timestamp,
        )
        if not ok:
            if text == "400":
                return is_succeeded
            await send_xml_to_ekinobilet(
                performance_data, show_data, payment, fond_kino_id,
                place, row, -price, payment_id, order_id,
            )
            await send_message_sync(
                chat_id,
                user_id,
                f"Извините, возникла ошибка, деньги вернутся на вашу карту.\n{text}",
            )
    except Exception:
        logger.exception("Ошибка валидации Почта Банк")

    return is_succeeded